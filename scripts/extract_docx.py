"""Extract plain text (paragraphs + tables) from a DOCX file for review."""
import sys

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_block_items(parent):
    from docx.oxml.ns import qn

    parent_elm = parent.element.body
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def main(path: str) -> None:
    doc = Document(path)
    print(f"# FILE: {path}")
    print(f"# paragraphs: {len(doc.paragraphs)}")
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                style = block.style.name if block.style is not None else ""
                print(f"[{style}] {text}")
        else:
            print("--- TABLE ---")
            for row in block.rows:
                cells = " | ".join(c.text.strip().replace("\n", " / ") for c in row.cells)
                print(f"  {cells}")
            print("--- END TABLE ---")
    # headers / footers
    for i, section in enumerate(doc.sections, 1):
        for part_name, part in (("header", section.header), ("footer", section.footer)):
            pieces = [p.text.strip() for p in part.paragraphs if p.text.strip()]
            if pieces:
                print(f"# {part_name} {i}: {' || '.join(pieces)}")


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        main(arg)
