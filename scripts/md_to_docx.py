"""把技术文档 Markdown 转换为排版良好的 Word（.docx）。

特性：标题层级、正文（中文宋体 / 西文 Times New Roman）、表格（Table Grid + 表题）、
Mermaid 图（若已由 scripts/render_mermaid.py 渲染为 PNG 则嵌入，否则保留代码块）、
代码块、引用、列表、公式行、静态目录、页脚页码、A4 页面。

用法::

    .venv\\Scripts\\python.exe scripts\\md_to_docx.py docs\\techdoc\\chapter4_method.md docs\\techdoc\\第四章_多模态多智能体协同预警方法.docx
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_BODY_FONT = "宋体"
CN_HEAD_FONT = "黑体"
EN_BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"

MERMAID_RE = re.compile(r"```mermaid\s*\n(.*?)```", re.S)

GREEK = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε", "varepsilon": "ε",
    "zeta": "ζ", "eta": "η", "theta": "θ", "iota": "ι", "kappa": "κ", "lambda": "λ",
    "mu": "μ", "nu": "ν", "xi": "ξ", "pi": "π", "rho": "ρ", "sigma": "σ", "tau": "τ",
    "upsilon": "υ", "phi": "φ", "varphi": "φ", "chi": "χ", "psi": "ψ", "omega": "ω",
    "Gamma": "Γ", "Delta": "Δ", "Theta": "Θ", "Lambda": "Λ", "Xi": "Ξ", "Pi": "Π",
    "Sigma": "Σ", "Phi": "Φ", "Psi": "Ψ", "Omega": "Ω",
}
MATH_SYMBOLS = {
    "in": "∈", "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥", "neq": "≠", "times": "×",
    "cdot": "·", "sum": "Σ", "prod": "Π", "partial": "∂", "infty": "∞", "approx": "≈",
    "rightarrow": "→", "to": "→", "Rightarrow": "⇒", "forall": "∀", "exists": "∃",
    "subset": "⊂", "subseteq": "⊆", "cup": "∪", "cap": "∩", "pm": "±", "mp": "∓",
}


def latex_to_text(text: str) -> str:
    """把行内/独立 LaTeX 公式转成可读 Unicode 文本（保留变量与关系符）。"""
    s = text.strip()
    s = re.sub(r"\\begin\{[^{}]*\}|\\end\{[^{}]*\}", "", s)
    s = re.sub(r"\\(?:hat|bar|tilde|overline)\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\(?:text|mathrm|mathcal|mathbb|mathbf|operatorname)\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\(?:left|right|,|;|!|quad|qquad|displaystyle)\b", "", s)
    for name, symbol in GREEK.items():
        s = re.sub(r"\\" + name + r"(?![A-Za-z])", symbol, s)
    for name, symbol in MATH_SYMBOLS.items():
        s = re.sub(r"\\" + name + r"(?![A-Za-z])", symbol, s)
    s = s.replace("\\{", "{").replace("\\}", "}").replace("\\%", "%").replace("\\\\", " ")
    s = re.sub(r"\\[A-Za-z]+", "", s)
    s = s.replace("\\", " ")
    s = s.replace("{", "").replace("}", "").replace("&", " ")
    s = s.replace("^", "")
    s = re.sub(r"\s+", " ", s).strip()
    return s


# --------------------------------------------------------------------------- #
# 低层工具
# --------------------------------------------------------------------------- #
def set_run_font(run, *, name: str = EN_BODY_FONT, east_asia: str = CN_BODY_FONT,
                 size: float | None = None, bold: bool = False, italic: bool = False,
                 color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline_runs(paragraph, text: str, *, base_size: float = 10.5, color: str | None = None) -> None:
    """解析 **加粗**、*斜体*、`代码`、$行内公式$ 并写入段落。"""
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`|\$[^$]+\$|\*[^*]+?\*)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, color=color)
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=CODE_FONT, east_asia=CODE_FONT, size=base_size - 0.5, color="8B1E3F")
        elif token.startswith("$") and token.endswith("$") and len(token) > 2:
            run = paragraph.add_run(latex_to_text(token[1:-1]))
            set_run_font(run, size=base_size, italic=True, color=color or "1F3864")
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size, italic=True, color=color)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=base_size, color=color)


def add_page_number_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = EN_BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_BODY_FONT)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(4)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[style_name]
        style.font.name = EN_BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_HEAD_FONT)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


# --------------------------------------------------------------------------- #
# 内容块
# --------------------------------------------------------------------------- #
def add_title_page(doc: Document, title: str) -> None:
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, name=CN_HEAD_FONT, east_asia=CN_HEAD_FONT, size=24, bold=True, color="1F3864")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("技术文档 · 论文方向2")
    set_run_font(run, size=14, color="595959")

    for _ in range(3):
        doc.add_paragraph()

    for line in (
        "跨境金融风险的多智能体协同预警：多模态信息融合视角",
        "版本：v0.4（与实现同步）    日期：2026-10-18",
        "系统实现：CBR-MAW 原型（六层两横切架构）",
        "",
        "说明：本文件为方法章技术文档，正文不包含实验结果；",
        "实验设计与结果分别见论文第五章、第六章。",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        set_run_font(run, size=11, color="404040")
    doc.add_page_break()


def add_toc(doc: Document, headings: list[tuple[int, str]]) -> None:
    p = doc.add_paragraph()
    run = p.add_run("目  录")
    set_run_font(run, name=CN_HEAD_FONT, east_asia=CN_HEAD_FONT, size=16, bold=True, color="1F3864")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    for level, text in headings:
        if level > 2:
            continue
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.left_indent = Cm(0 if level == 1 else 0.8)
        run = para.add_run(text)
        set_run_font(run, size=11 if level == 1 else 10.5, bold=(level == 1))
    doc.add_page_break()


def add_table(doc: Document, rows: list[list[str]], caption: str | None) -> None:
    if not rows:
        return
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        set_run_font(run, size=9.5, bold=True, color="1F3864")
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(2)
        cap.paragraph_format.keep_with_next = True
    n_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(n_cols):
            cell = table.cell(r_idx, c_idx)
            text = row[c_idx] if c_idx < len(row) else ""
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.line_spacing = 1.15
            add_inline_runs(para, text.strip(), base_size=8.5)
            if r_idx == 0:
                for run in para.runs:
                    run.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, code: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.4)
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.line_spacing = 1.1
    run = para.add_run(code.rstrip("\n"))
    set_run_font(run, name=CODE_FONT, east_asia=CODE_FONT, size=8.5, color="333333")
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    para._p.get_or_add_pPr().append(shading)


def add_figure(doc: Document, image_path: Path, caption: str | None, width_cm: float = 16.0) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        set_run_font(run, size=9.5, bold=True, color="1F3864")
        cap.paragraph_format.space_after = Pt(8)


# --------------------------------------------------------------------------- #
# Markdown 解析
# --------------------------------------------------------------------------- #
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?$")


def split_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def convert(md_path: Path, out_path: Path, figure_dir: Path | None = None) -> None:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    figure_dir = figure_dir or md_path.parent / "figures"
    figure_files = sorted(figure_dir.glob("fig_*.png")) if figure_dir.exists() else []
    figure_cursor = 0

    headings: list[tuple[int, str]] = []
    for line in lines:
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m and not line.startswith("```"):
            headings.append((len(m.group(1)), m.group(2).strip()))

    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Cm(2.3 if attr in ("left_margin", "right_margin") else 2.2))
    setup_styles(doc)
    add_page_number_footer(section)

    title = headings[0][1] if headings and headings[0][0] == 1 else md_path.stem
    add_title_page(doc, title)
    add_toc(doc, [h for h in headings if h[0] <= 3 and h[1] != title])

    i = 0
    pending_caption: str | None = None
    n_tables = 0
    n_figures = 0
    n_code = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            lang = stripped[3:].strip().lower()
            block: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            if lang == "mermaid":
                if figure_cursor < len(figure_files):
                    add_figure(doc, figure_files[figure_cursor], pending_caption)
                    pending_caption = None
                    figure_cursor += 1
                    n_figures += 1
                else:
                    add_code_block(doc, "\n".join(block))
                    n_code += 1
            else:
                add_code_block(doc, "\n".join(block))
                n_code += 1
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1].strip()):
            rows: list[list[str]] = [split_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                if TABLE_SEP_RE.match(lines[i].strip()):
                    pass
                else:
                    rows.append(split_table_row(lines[i]))
                i += 1
            add_table(doc, rows, pending_caption)
            pending_caption = None
            n_tables += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            heading_text = m.group(2).strip()
            if level == 1:
                para = doc.add_paragraph(style="Heading 1")
                add_inline_runs(para, heading_text, base_size=16)
            elif level == 2:
                para = doc.add_paragraph(style="Heading 2")
                add_inline_runs(para, heading_text, base_size=13)
            else:
                para = doc.add_paragraph(style="Heading 3")
                add_inline_runs(para, heading_text, base_size=11.5)
            i += 1
            continue

        # 图/表题注（单独一行）
        if re.match(r"^(图|表)\s*\d+[-–]\d+", stripped):
            if stripped.startswith("图") and figure_cursor < len(figure_files) and pending_caption is None:
                # 图片已在前一个 mermaid 块插入，这里作为图题
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(stripped)
                set_run_font(run, size=9.5, bold=True, color="1F3864")
                cap.paragraph_format.space_after = Pt(8)
            else:
                pending_caption = stripped
            i += 1
            continue

        # 公式（$$ ... $$，可跨行）
        if stripped.startswith("$$"):
            formula_lines: list[str] = []
            first = stripped[2:]
            if first.strip().endswith("$$") and first.strip() != "$$":
                formula_lines.append(first.strip()[:-2])
            else:
                if first.strip():
                    formula_lines.append(first.strip())
                i += 1
                while i < len(lines) and "$$" not in lines[i]:
                    formula_lines.append(lines[i].strip())
                    i += 1
                if i < len(lines):
                    tail = lines[i].strip()
                    formula_lines.append(tail.replace("$$", "").strip())
            formula_text = latex_to_text(" ".join(x for x in formula_lines if x))
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(formula_text)
            set_run_font(run, size=10.5, italic=True, color="1F3864")
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.6)
            add_inline_runs(para, stripped.lstrip("> ").strip(), base_size=9.5, color="595959")
            i += 1
            continue

        # 列表
        if re.match(r"^[-*+]\s+", stripped):
            para = doc.add_paragraph(style="List Bullet")
            add_inline_runs(para, re.sub(r"^[-*+]\s+", "", stripped))
            i += 1
            continue
        if re.match(r"^\d+[.、)]\s+", stripped):
            para = doc.add_paragraph(style="List Number")
            add_inline_runs(para, re.sub(r"^\d+[.、)]\s+", "", stripped))
            i += 1
            continue

        # 普通段落
        para = doc.add_paragraph()
        add_inline_runs(para, stripped)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"[OK] DOCX: {out_path}")
    print(f"     标题 {len(headings)} 个 / 表格 {n_tables} 张 / 图 {n_figures} 张 / 代码块 {n_code} 个")


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print("用法: python scripts/md_to_docx.py <input.md> <output.docx> [figures_dir]")
        return 2
    convert(Path(argv[0]), Path(argv[1]), Path(argv[2]) if len(argv) > 2 else None)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
